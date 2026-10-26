"""本机格式转换组件；仅通过固定 JSON 协议处理已授权的临时文件。"""

from contextlib import contextmanager
import csv
import io
import json
import math
import os
from pathlib import Path
import re
import stat
import subprocess
import sys
import tempfile
import time

VERSION = "1.0.0"
INPUT_LIMIT = 200 * 1024 * 1024
OUTPUT_LIMIT = 2 * 1024 * 1024 * 1024
PIXEL_LIMIT = 50_000_000
PAGE_LIMIT = 500
TEXT_LIMIT = 32 * 1024 * 1024
TARGETS = {"png", "jpg", "pdf", "docx", "txt", "xlsx", "csv", "pptx"}
IMAGES = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tif", ".tiff", ".gif", ".ico"}


class ConversionError(Exception):
    pass


def checked_path(value, directory=False):
    if not isinstance(value, str) or not value or "\x00" in value:
        raise ConversionError("Conversion invalid path")
    path = Path(value)
    if not path.is_absolute() or value.startswith(("\\\\", "//")):
        raise ConversionError("Conversion requires local absolute paths")
    for part in [path, *path.parents]:
        if part.is_symlink() or getattr(part.stat(), "st_file_attributes", 0) & 0x400:
            raise ConversionError("Conversion rejects reparse points")
    details = path.stat()
    if directory:
        if not stat.S_ISDIR(details.st_mode):
            raise ConversionError("Conversion output directory unavailable")
    elif not stat.S_ISREG(details.st_mode) or details.st_size == 0 or details.st_size > INPUT_LIMIT:
        raise ConversionError("Conversion input exceeds limit or is not a regular file")
    return path


def page_indices(selection, count, explicit=False):
    if count < 1 or count > PAGE_LIMIT:
        raise ConversionError("Conversion page limit exceeded")
    if not isinstance(selection, str) or len(selection) > 2048:
        raise ConversionError("Conversion invalid pages")
    if not selection.strip() or selection.strip().lower() == "all":
        if explicit and count > 1:
            raise ConversionError("Conversion requires an explicit frame/page selection")
        return list(range(count))
    result = []
    for token in selection.split(","):
        match = re.fullmatch(r"\s*([1-9][0-9]{0,3})(?:\s*-\s*([1-9][0-9]{0,3}))?\s*", token)
        if not match:
            raise ConversionError("Conversion invalid pages")
        first, last = int(match[1]), int(match[2] or match[1])
        if first > last or last > count:
            raise ConversionError("Conversion pages out of range")
        for number in range(first - 1, last):
            if number not in result:
                result.append(number)
    return result


class OutputBudget:
    def __init__(self, root):
        self.root = root
        self.paths = []
        self.bytes = 0

    @contextmanager
    def create(self, name):
        path = self.root / name
        stream = path.open("xb")
        try:
            yield BudgetFile(stream, OUTPUT_LIMIT - self.bytes)
            stream.flush()
            size = stream.seek(0, io.SEEK_END)
            if size < 1:
                raise ConversionError("Conversion produced empty output")
            self.bytes += size
            self.paths.append(path)
        except BaseException:
            stream.close()
            path.unlink(missing_ok=True)
            raise
        finally:
            stream.close()

    def cleanup(self):
        for path in self.paths:
            path.unlink(missing_ok=True)


class BudgetFile:
    def __init__(self, stream, maximum):
        self.stream = stream
        self.maximum = maximum

    def write(self, data):
        if self.stream.tell() + len(data) > self.maximum:
            raise ConversionError("Conversion output exceeds limit")
        return self.stream.write(data)

    def __getattr__(self, name):
        return getattr(self.stream, name)


def check_pixels(width, height):
    if width <= 0 or height <= 0 or width * height > PIXEL_LIMIT:
        raise ConversionError("Conversion pixel limit exceeded")


def white_image(image):
    from PIL import Image, ImageOps
    check_pixels(*image.size)
    image = ImageOps.exif_transpose(image)
    rgba = image.convert("RGBA")
    result = Image.new("RGB", rgba.size, "white")
    result.paste(rgba, mask=rgba.getchannel("A"))
    return result


def selected_images(source, pages):
    from PIL import Image
    Image.MAX_IMAGE_PIXELS = PIXEL_LIMIT
    with Image.open(source) as picture:
        indices = page_indices(pages, getattr(picture, "n_frames", 1), explicit=True)
        for index in indices:
            picture.seek(index)
            yield white_image(picture)


def image_to_pdf(sources, options, budget, work):
    from pypdf import PdfReader, PdfWriter
    writer = PdfWriter()
    count = 0
    retained = 0
    for source in sources:
        for picture in selected_images(source, options["pages"]):
            count += 1
            if count > PAGE_LIMIT:
                raise ConversionError("Conversion page limit exceeded")
            # 逐张转为有界页面，避免把所有原始像素同时留在内存中。
            single = work / f"image-{count:04}.pdf"
            with single.open("xb") as stream:
                picture.save(BudgetFile(stream, OUTPUT_LIMIT - budget.bytes), format="PDF", resolution=options["dpi"], quality=options["quality"])
            picture.close()
            reader = PdfReader(single)
            writer.add_page(reader.pages[0])
            retained += single.stat().st_size
            single.unlink()
            if retained > OUTPUT_LIMIT:
                raise ConversionError("Conversion output exceeds limit")
    with budget.create("converted.pdf") as output:
        writer.write(output)


@contextmanager
def pdf_document(source):
    import pypdfium2 as pdfium
    document = pdfium.PdfDocument(source)
    try:
        if len(document) > PAGE_LIMIT:
            raise ConversionError("Conversion page limit exceeded")
        yield document
    finally:
        document.close()


def rendered_pages(source, options):
    with pdf_document(source) as document:
        for index in page_indices(options["pages"], len(document)):
            page = document[index]
            try:
                width, height = page.get_size()
                if not math.isfinite(width) or not math.isfinite(height):
                    raise ConversionError("Conversion invalid PDF geometry")
                scale = options["dpi"] / 72
                check_pixels(math.ceil(width * scale), math.ceil(height * scale))
                bitmap = page.render(scale=scale, fill_color=(255, 255, 255, 255))
                try:
                    image = bitmap.to_pil().convert("RGB")
                finally:
                    bitmap.close()
                yield index, (width, height), image
            finally:
                page.close()


def pdf_images(source, options, budget):
    target = options["target"]
    for index, _, picture in rendered_pages(source, options):
        with budget.create(f"page-{index + 1:04}.{target}") as output:
            picture.save(output, format="PNG" if target == "png" else "JPEG", quality=options["quality"])
        picture.close()


def pdf_text(source, options):
    import pdfplumber
    pages = []
    total = 0
    with pdfplumber.open(source) as document:
        for index in page_indices(options["pages"], len(document.pages)):
            page = document.pages[index]
            try:
                text = page.extract_text() or ""
                total += len(text.encode("utf-8"))
                if total > TEXT_LIMIT:
                    raise ConversionError("Conversion extracted text exceeds limit")
                pages.append(text)
            finally:
                page.close()
    if not any(text.strip() for text in pages):
        raise ConversionError("Conversion PDF has no extractable text; install OCR component and enable OCR")
    return pages


def write_text(pages, target, budget):
    if target == "txt":
        with budget.create("converted.txt") as output:
            output.write("\n\n\f\n\n".join(pages).encode("utf-8"))
        return
    from docx import Document
    from docx.shared import Pt
    document = Document()
    document.styles["Normal"].font.name = "Segoe UI"
    document.styles["Normal"].font.size = Pt(11)
    for index, text in enumerate(pages):
        if index:
            document.add_page_break()
        for line in text.splitlines():
            # OOXML 不接受 XML 控制字符；正文始终作为文字写入。
            document.add_paragraph(re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", line))
    with budget.create("converted.docx") as output:
        document.save(output)


def neutral_csv(value):
    text = str(value or "")
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r", "\n")) else text


def pdf_tables(source, options, budget):
    import pdfplumber
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.remove(workbook.active)
    found = 0
    cells = 0
    with pdfplumber.open(source) as document:
        for index in page_indices(options["pages"], len(document.pages)):
            page = document.pages[index]
            try:
                tables = page.extract_tables()
                for table in tables:
                    if not table or not any(any(cell for cell in row) for row in table):
                        continue
                    found += 1
                    cells += sum(len(row) for row in table)
                    if found > 500 or cells > 1_000_000:
                        raise ConversionError("Conversion table limits exceeded")
                    if options["target"] == "csv":
                        text = io.StringIO(newline="")
                        writer = csv.writer(text)
                        writer.writerows([[neutral_csv(value) for value in row] for row in table])
                        with budget.create(f"table-{found:04}.csv") as output:
                            output.write(text.getvalue().encode("utf-8-sig"))
                    else:
                        sheet = workbook.create_sheet(f"Page {index + 1} Table {found}")
                        for row_index, row in enumerate(table, 1):
                            for column_index, value in enumerate(row, 1):
                                cell = sheet.cell(row_index, column_index, str(value or ""))
                                # 强制字符串，绝不把从 PDF 提取的内容升级成公式。
                                cell.data_type = "s"
            finally:
                page.close()
    if found == 0:
        raise ConversionError("Conversion No tables found in PDF")
    if options["target"] == "xlsx":
        with budget.create("converted.xlsx") as output:
            workbook.save(output)
    workbook.close()


def pdf_presentation(source, options, budget, work):
    from pptx import Presentation
    from pptx.util import Inches
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    for number, (_, dimensions, picture) in enumerate(rendered_pages(source, options)):
        if number == 0:
            width, height = dimensions
            ratio = min(13.333333 / (width / 72), 7.5 / (height / 72))
            presentation.slide_width = Inches(width / 72 * ratio)
            presentation.slide_height = Inches(height / 72 * ratio)
        path = work / f"slide-{number:04}.png"
        with path.open("xb") as stream:
            picture.save(BudgetFile(stream, OUTPUT_LIMIT - budget.bytes), format="PNG")
        picture.close()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        width, height = dimensions
        scale = min(presentation.slide_width / width, presentation.slide_height / height)
        target_width, target_height = int(width * scale), int(height * scale)
        slide.shapes.add_picture(str(path), (presentation.slide_width - target_width) // 2, (presentation.slide_height - target_height) // 2, target_width, target_height)
        path.unlink()
    with budget.create("converted.pptx") as output:
        presentation.save(output)


def workspace_bytes(root):
    total = 0
    for path in root.rglob("*"):
        if path.is_symlink():
            raise ConversionError("Conversion rejects temporary links")
        if path.is_file():
            total += path.stat().st_size
            if total > OUTPUT_LIMIT:
                raise ConversionError("Conversion workspace exceeds limit")
    return total


def run_ocr(command, environment, root):
    # 引擎写出的文件也计入工作区；不把不可信文档的诊断输出积存在内存中。
    child = subprocess.Popen(command, stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=environment, creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
    deadline = time.monotonic() + 120
    try:
        while child.poll() is None:
            workspace_bytes(root)
            if time.monotonic() > deadline:
                raise ConversionError("Conversion OCR timed out")
            time.sleep(0.05)
        workspace_bytes(root)
        if child.returncode:
            raise ConversionError("Conversion OCR engine failed")
    finally:
        if child.poll() is None:
            child.kill()
        child.wait(timeout=10)


def ocr_pages(source, options, request, budget, work):
    from pypdf import PdfReader, PdfWriter
    executable_value, data_value = request.get("ocrExecutable"), request.get("ocrData")
    if not executable_value or not data_value:
        raise ConversionError("Conversion OCR component is not installed")
    executable = checked_path(executable_value)
    data = checked_path(data_value, directory=True)
    if executable.name.lower() != "tesseract.exe" or not all((data / f"{language}.traineddata").is_file() for language in ("chi_sim", "eng")):
        raise ConversionError("Conversion OCR component or language data is unavailable")
    images = ((index, picture) for index, _, picture in rendered_pages(source, options)) if source.suffix.lower() == ".pdf" else enumerate(selected_images(source, options["pages"]))
    pages, writer = [], PdfWriter()
    for index, picture in images:
        image_path, output_base = work / f"ocr-{index:04}.png", work / f"recognized-{index:04}"
        with image_path.open("xb") as image_output:
            picture.save(BudgetFile(image_output, OUTPUT_LIMIT - budget.bytes), format="PNG", dpi=(options["dpi"], options["dpi"]))
        picture.close()
        command = [str(executable), str(image_path), str(output_base), "--tessdata-dir", str(data), "-l", "chi_sim+eng", "--dpi", str(options["dpi"]), "pdf" if options["target"] == "pdf" else "txt"]
        environment = dict(os.environ)
        environment["OMP_THREAD_LIMIT"] = "1"
        environment.pop("TESSDATA_PREFIX", None)
        run_ocr(command, environment, budget.root)
        image_path.unlink()
        if options["target"] == "pdf":
            page_output = output_base.with_suffix(".pdf")
            writer.append(PdfReader(page_output), import_outline=False)
            page_output.unlink()
        else:
            page_output = output_base.with_suffix(".txt")
            if page_output.stat().st_size + sum(len(text.encode("utf-8")) for text in pages) > TEXT_LIMIT:
                raise ConversionError("Conversion extracted text exceeds limit")
            pages.append(page_output.read_text(encoding="utf-8"))
            page_output.unlink()
    if options["target"] == "pdf":
        with budget.create("searchable.pdf") as output:
            writer.write(output)
    else:
        if not any(text.strip() for text in pages):
            raise ConversionError("Conversion OCR found no readable text")
        write_text(pages, options["target"], budget)


def convert(request):
    if not isinstance(request, dict) or set(request) - {"source", "sources", "outputDir", "options", "ocrExecutable", "ocrData"}:
        raise ConversionError("Conversion invalid request")
    source = checked_path(request.get("source"))
    output = checked_path(request.get("outputDir"), directory=True)
    options = {"quality": 90, "dpi": 150, "pages": "", "ocr": False, "merge": False, **request.get("options", {})}
    target = options.get("target")
    if target not in TARGETS or type(options["quality"]) is not int or not 1 <= options["quality"] <= 100 or options["dpi"] not in (150, 300) or type(options["ocr"]) is not bool or type(options["merge"]) is not bool:
        raise ConversionError("Conversion invalid target or options")
    suffix = source.suffix.lower()
    if suffix not in IMAGES | {".pdf"}:
        raise ConversionError("Conversion unsupported helper input")
    sources = [source]
    if "sources" in request:
        if not isinstance(request["sources"], list) or not 1 <= len(request["sources"]) <= 100 or not options["merge"] or target != "pdf" or options["ocr"]:
            raise ConversionError("Conversion invalid image merge")
        sources = [checked_path(value) for value in request["sources"]]
        if any(path.suffix.lower() not in IMAGES for path in sources):
            raise ConversionError("Conversion merge requires images")
    budget = OutputBudget(output)
    try:
        with tempfile.TemporaryDirectory(prefix=".helper-", dir=output) as temporary:
            work = Path(temporary)
            if options["ocr"]:
                if target not in {"txt", "docx", "pdf"}:
                    raise ConversionError("Conversion OCR supports TXT, DOCX or searchable PDF")
                ocr_pages(source, options, request, budget, work)
            elif suffix in IMAGES and target == "pdf":
                image_to_pdf(sources, options, budget, work)
            elif suffix == ".pdf":
                if target in {"png", "jpg"}:
                    pdf_images(source, options, budget)
                elif target in {"txt", "docx"}:
                    write_text(pdf_text(source, options), target, budget)
                elif target in {"xlsx", "csv"}:
                    pdf_tables(source, options, budget)
                elif target == "pptx":
                    pdf_presentation(source, options, budget, work)
                else:
                    raise ConversionError("Conversion unsupported PDF direction")
            else:
                raise ConversionError("Conversion unsupported helper direction")
        return {"outputs": [str(path) for path in budget.paths]}
    except BaseException:
        budget.cleanup()
        raise


def main():
    if sys.argv[1:] in (["--version"], ["--probe"]):
        print(f"SuperSearch conversion helper {VERSION}")
        return 0
    if sys.argv[1:]:
        print(json.dumps({"outputs": [], "error": "Conversion invalid arguments"}))
        return 1
    try:
        raw = sys.stdin.buffer.read(128 * 1024 + 1)
        if len(raw) > 128 * 1024:
            raise ConversionError("Conversion request exceeds limit")
        response = convert(json.loads(raw))
        print(json.dumps(response, ensure_ascii=True))
        return 0
    except ConversionError as error:
        message = str(error)
    except subprocess.TimeoutExpired:
        message = "Conversion OCR timed out"
    except FileExistsError:
        message = "Conversion output already exists"
    except (OSError, ValueError, TypeError, KeyError, json.JSONDecodeError):
        message = "Conversion invalid file or request"
    except Exception:
        # 不向协议或日志输出文档正文、完整路径和库内部回溯。
        message = "Conversion document processing failed"
    print(json.dumps({"outputs": [], "error": message}, ensure_ascii=True))
    return 1


if __name__ == "__main__":
    raise SystemExit(main())
