"""用真实文件验证本地 PDF 转换组件。"""

import csv
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import importlib.util
import random

from PIL import Image
from docx import Document
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader, PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject


HERE = Path(__file__).resolve().parent
HELPER = HERE / "helper.py"


def make_pdf(path, table=False, pages=1, text="Hello local conversion"):
    writer = PdfWriter()
    for index in range(pages):
        page = writer.add_blank_page(width=320, height=240)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        content = f"BT /F1 12 Tf 30 210 Td ({text} {index + 1}) Tj ET\n"
        if table:
            content += "0.5 w 30 80 m 290 80 l 290 180 l 30 180 l h S 160 80 m 160 180 l S 30 130 m 290 130 l S\n"
            content += "BT /F1 12 Tf 40 150 Td (Name) Tj 130 0 Td (Value) Tj ET\n"
            content += "BT /F1 12 Tf 40 100 Td (Budget) Tj 130 0 Td (=1+2) Tj ET\n"
        stream = DecodedStreamObject()
        stream.set_data(content.encode("ascii"))
        page[NameObject("/Contents")] = writer._add_object(stream)
    with open(path, "wb") as output:
        writer.write(output)


class HelperTests(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix="conversion-helper-")
        self.root = Path(self.temporary.name)
        self.source = self.root / "sample.pdf"
        self.output = self.root / "out"
        self.output.mkdir()
        make_pdf(self.source)

    def tearDown(self):
        self.temporary.cleanup()

    def invoke(self, target, source=None, **options):
        request = {"source": str(source or self.source), "outputDir": str(self.output), "options": {"target": target, "quality": 90, "dpi": 150, "icoSize": 256, "pages": "", "sheet": 1, "merge": False, "ocr": False, **options}}
        if "sources" in options:
            request["sources"] = options.pop("sources")
            request["options"].pop("sources", None)
        if "ocrExecutable" in options:
            request["ocrExecutable"] = options["ocrExecutable"]
            request["options"].pop("ocrExecutable", None)
        if "ocrData" in options:
            request["ocrData"] = options["ocrData"]
            request["options"].pop("ocrData", None)
        result = subprocess.run([sys.executable, str(HELPER)], input=json.dumps(request), text=True, capture_output=True, timeout=45, encoding="utf-8")
        self.assertTrue(result.stdout.strip(), result.stderr)
        response = json.loads(result.stdout)
        return result, response

    def successful(self, target, **options):
        result, response = self.invoke(target, **options)
        self.assertEqual(result.returncode, 0, response)
        self.assertNotIn("error", response)
        paths = [Path(value) for value in response["outputs"]]
        self.assertTrue(paths)
        for path in paths:
            self.assertTrue(path.is_relative_to(self.output))
            self.assertGreater(path.stat().st_size, 0)
        return paths

    def test_probe_does_not_wait_for_stdin(self):
        result = subprocess.run([sys.executable, str(HELPER), "--version"], capture_output=True, text=True, timeout=10)
        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertIn("1.0.0", result.stdout)

    def test_text_pdf_produces_editable_word_paragraphs(self):
        paths = self.successful("docx")
        self.assertIn("Hello local conversion", "\n".join(p.text for p in Document(paths[0]).paragraphs))

    def test_text_pdf_produces_plain_utf8(self):
        paths = self.successful("txt")
        self.assertIn("Hello local conversion", paths[0].read_text(encoding="utf-8"))

    def test_table_pdf_exports_xlsx_without_formula_execution(self):
        make_pdf(self.source, table=True)
        paths = self.successful("xlsx")
        book = load_workbook(paths[0], data_only=False)
        self.assertEqual(book.active["A1"].value, "Name")
        self.assertEqual(book.active["B2"].value, "=1+2")
        self.assertEqual(book.active["B2"].data_type, "s")

    def test_csv_table_export_neutralizes_formula_prefixes(self):
        make_pdf(self.source, table=True)
        paths = self.successful("csv")
        with paths[0].open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.reader(stream))
        self.assertEqual(rows[0], ["Name", "Value"])
        self.assertEqual(rows[1][1], "'=1+2")

    def test_pdf_without_tables_does_not_report_empty_success(self):
        result, response = self.invoke("xlsx")
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("No tables", response["error"])
        self.assertEqual(list(self.output.iterdir()), [])

    def test_pdf_renders_only_requested_pages(self):
        make_pdf(self.source, pages=3)
        paths = self.successful("png", pages="2-3")
        self.assertEqual(len(paths), 2)
        with Image.open(paths[0]) as picture:
            self.assertGreater(picture.width, 600)
            self.assertEqual(picture.format, "PNG")

    def test_pdf_pptx_has_one_picture_per_page(self):
        make_pdf(self.source, pages=2)
        paths = self.successful("pptx")
        slides = Presentation(paths[0])
        self.assertEqual(len(slides.slides), 2)
        self.assertTrue(all(len(slide.shapes) == 1 for slide in slides.slides))

    def test_images_merge_into_ordered_pdf_with_white_background(self):
        first, second = self.root / "one.png", self.root / "two.png"
        Image.new("RGBA", (120, 60), (255, 0, 0, 0)).save(first)
        Image.new("RGB", (60, 120), "blue").save(second)
        paths = self.successful("pdf", source=first, merge=True, sources=[str(first), str(second)])
        self.assertEqual(len(PdfReader(paths[0]).pages), 2)

    def test_animation_requires_an_explicit_frame(self):
        source = self.root / "animated.gif"
        one, two = Image.new("RGB", (20, 20), "red"), Image.new("RGB", (20, 20), "blue")
        one.save(source, save_all=True, append_images=[two], duration=100, loop=0)
        result, response = self.invoke("pdf", source=source)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("frame", response["error"])
        self.assertEqual(len(PdfReader(self.successful("pdf", source=source, pages="2")[0]).pages), 1)

    def test_blank_scan_does_not_generate_empty_word(self):
        scan = self.root / "scan.png"
        Image.new("RGB", (60, 60), "white").save(scan)
        scanned_pdf = self.root / "scan.pdf"
        Image.open(scan).save(scanned_pdf, "PDF")
        result, response = self.invoke("docx", source=scanned_pdf)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("OCR", response["error"])

    def test_unknown_target_and_bad_pages_fail_before_writing(self):
        for target, options in [("exe", {}), ("png", {"pages": "0"}), ("png", {"pages": "2-1"}), ("png", {"pages": "999"})]:
            result, response = self.invoke(target, **options)
            self.assertNotEqual(result.returncode, 0, response)
            self.assertEqual(list(self.output.iterdir()), [])

    def test_live_working_files_stay_within_output_limit(self):
        spec = importlib.util.spec_from_file_location("helper_under_test", HERE / "helper.py")
        module = importlib.util.module_from_spec(spec)
        spec.loader.exec_module(module)
        first, second = self.root / "noise-one.png", self.root / "noise-two.png"
        noise = random.Random(17)
        for path in (first, second):
            Image.frombytes("RGB", (300, 300), noise.randbytes(300 * 300 * 3)).save(path)
        maximum = 0
        original = module.BudgetFile.write

        def observed_write(stream, data):
            nonlocal maximum
            result = original(stream, data)
            stream.stream.flush()
            maximum = max(maximum, sum(path.stat().st_size for path in self.output.rglob("*") if path.is_file()))
            return result

        with patch.object(module, "OUTPUT_LIMIT", 100_000), patch.object(module.BudgetFile, "write", observed_write):
            with self.assertRaises(module.ConversionError):
                module.convert({"source": str(first), "sources": [str(first), str(second)], "outputDir": str(self.output), "options": {"target": "pdf", "merge": True}})
        self.assertLessEqual(maximum, 100_000)
        self.assertEqual(list(self.output.iterdir()), [])

    def test_ocr_requires_real_engine_and_language_data(self):
        result, response = self.invoke("txt", ocr=True)
        self.assertNotEqual(result.returncode, 0)
        self.assertIn("OCR component", response["error"])

    @unittest.skipUnless(os.environ.get("SUPERSEARCH_TEST_TESSERACT"), "Set SUPERSEARCH_TEST_TESSERACT and SUPERSEARCH_TEST_TESSDATA for real OCR")
    def test_real_ocr_generates_text_and_searchable_pdf(self):
        from PIL import ImageDraw, ImageFont
        image_path = self.root / "ocr.png"
        picture = Image.new("RGB", (1000, 160), "white")
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 60)
        ImageDraw.Draw(picture).text((30, 40), "LOCAL CONVERSION 12345", font=font, fill="black")
        picture.save(image_path)
        options = {"source": image_path, "ocr": True, "ocrExecutable": os.environ["SUPERSEARCH_TEST_TESSERACT"], "ocrData": os.environ["SUPERSEARCH_TEST_TESSDATA"]}
        text = self.successful("txt", **options)[0].read_text(encoding="utf-8")
        self.assertIn("CONVERSION", text)
        result = self.successful("pdf", **options)[0]
        self.assertIn("CONVERSION", PdfReader(result).pages[0].extract_text())

    @unittest.skipUnless(os.environ.get("SUPERSEARCH_TEST_TESSERACT"), "Set OCR component paths for Chinese scan verification")
    def test_real_chinese_scan_produces_editable_word(self):
        from PIL import ImageDraw, ImageFont
        picture = Image.new("RGB", (1500, 240), "white")
        font = ImageFont.truetype("C:/Windows/Fonts/simsun.ttc", 72)
        ImageDraw.Draw(picture).text((40, 65), "本地格式转换 测试 12345", font=font, fill="black")
        scan = self.root / "chinese-scan.pdf"
        picture.save(scan, format="PDF", resolution=150)
        result = self.successful("docx", source=scan, ocr=True, ocrExecutable=os.environ["SUPERSEARCH_TEST_TESSERACT"], ocrData=os.environ["SUPERSEARCH_TEST_TESSDATA"])[0]
        text = "\n".join(paragraph.text for paragraph in Document(result).paragraphs)
        self.assertIn("本地", text)
        self.assertIn("转换", text)
        self.assertIn("12345", text)


if __name__ == "__main__":
    unittest.main()
